from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, or_, func
from ..database import get_private_db
from ..models.transaction import Transaction
from ..models.account import Account
from ..models.category import Category
from ..schemas.transaction import (
    TransactionCreate, TransactionUpdate, TransactionResponse,
    ImportPreviewRow, ImportPreviewResponse, ImportCommitRequest, ImportResult
)
from ..middleware.auth import get_current_user_id
from typing import List
import io
import csv
import re

router = APIRouter(prefix="/api/transactions", tags=["transactions"])


@router.get("", response_model=List[TransactionResponse])
async def get_transactions(
    year: int | None = None,
    month: int | None = None,
    type_filter: str | None = None,
    account_id: str | None = None,
    page: int = Query(1, ge=1),
    page_size: int = Query(50, ge=1, le=500),
    user_id: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_private_db)
):
    query = select(Transaction).where(Transaction.user_id == user_id)

    if year and month:
        month_str = f"{year}-{month:02d}"
        query = query.where(Transaction.date.like(f"{month_str}%"))
    if type_filter:
        query = query.where(Transaction.type == type_filter)
    if account_id:
        query = query.where(or_(Transaction.account_id == account_id, Transaction.dest_account_id == account_id))

    query = query.order_by(Transaction.date.desc(), Transaction.created_at.desc())

    # 按月查询时返回该月全部交易（记账页需要完整月数据做统计/分组）；
    # 无年月过滤时分页（用于通用列表/历史浏览）。
    if year and month:
        result = await db.execute(query)
    else:
        query = query.offset((page - 1) * page_size).limit(page_size)
        result = await db.execute(query)
    txns = result.scalars().all()

    account_ids = set(t.account_id for t in txns if t.account_id) | set(t.dest_account_id for t in txns if t.dest_account_id)
    cat_ids = set(t.category_id for t in txns if t.category_id)

    acc_result = await db.execute(select(Account).where(Account.id.in_(account_ids)))
    acc_map = {a.id: a.name for a in acc_result.scalars().all()}

    cat_result = await db.execute(select(Category).where(Category.id.in_(cat_ids)))
    cat_map = {c.id: (c.name, c.color) for c in cat_result.scalars().all()}

    responses = []
    for t in txns:
        cat_name, cat_color = cat_map.get(t.category_id, ("", "")) if t.category_id else ("", "")
        responses.append(TransactionResponse(
            id=t.id, user_id=t.user_id, type=t.type, date=t.date, amount=t.amount,
            account_id=t.account_id, dest_account_id=t.dest_account_id, category_id=t.category_id,
            tag_ids=t.tag_ids or [], description=t.description, remark=t.remark,
            location=t.location or "",
            created_at=str(t.created_at), updated_at=str(t.updated_at),
            account_name=acc_map.get(t.account_id, ""),
            category_name=cat_name, category_color=cat_color
        ))
    return responses


@router.post("", response_model=TransactionResponse)
async def create_transaction(req: TransactionCreate, user_id: str = Depends(get_current_user_id), db: AsyncSession = Depends(get_private_db)):
    acc_result = await db.execute(select(Account).where(Account.id == req.account_id, Account.user_id == user_id))
    account = acc_result.scalar_one_or_none()
    if not account:
        raise HTTPException(status_code=400, detail="Invalid account")

    txn = Transaction(user_id=user_id, **req.model_dump())
    db.add(txn)
    await db.commit()
    await db.refresh(txn)

    cat_name, cat_color = "", ""
    if txn.category_id:
        cat_result = await db.execute(select(Category).where(Category.id == txn.category_id))
        cat = cat_result.scalar_one_or_none()
        if cat:
            cat_name, cat_color = cat.name, cat.color

    return TransactionResponse(
        id=txn.id, user_id=txn.user_id, type=txn.type, date=txn.date, amount=txn.amount,
        account_id=txn.account_id, dest_account_id=txn.dest_account_id, category_id=txn.category_id,
        tag_ids=txn.tag_ids or [], description=txn.description, remark=txn.remark,
        location=txn.location or "",
        created_at=str(txn.created_at), updated_at=str(txn.updated_at),
        account_name=account.name, category_name=cat_name, category_color=cat_color
    )


@router.put("/{transaction_id}", response_model=TransactionResponse)
async def update_transaction(transaction_id: str, req: TransactionUpdate, user_id: str = Depends(get_current_user_id), db: AsyncSession = Depends(get_private_db)):
    result = await db.execute(select(Transaction).where(Transaction.id == transaction_id, Transaction.user_id == user_id))
    txn = result.scalar_one_or_none()
    if not txn:
        raise HTTPException(status_code=404, detail="Transaction not found")

    for key, value in req.model_dump(exclude_unset=True).items():
        setattr(txn, key, value)
    await db.commit()
    await db.refresh(txn)

    return TransactionResponse(
        id=txn.id, user_id=txn.user_id, type=txn.type, date=txn.date, amount=txn.amount,
        account_id=txn.account_id, dest_account_id=txn.dest_account_id, category_id=txn.category_id,
        tag_ids=txn.tag_ids or [], description=txn.description, remark=txn.remark,
        location=txn.location or "",
        created_at=str(txn.created_at), updated_at=str(txn.updated_at),
        account_name="", category_name="", category_color=""
    )


@router.delete("/{transaction_id}")
async def delete_transaction(transaction_id: str, user_id: str = Depends(get_current_user_id), db: AsyncSession = Depends(get_private_db)):
    result = await db.execute(select(Transaction).where(Transaction.id == transaction_id, Transaction.user_id == user_id))
    txn = result.scalar_one_or_none()
    if not txn:
        raise HTTPException(status_code=404, detail="Transaction not found")
    await db.delete(txn)
    await db.commit()
    return {"message": "Transaction deleted"}


# =====================================================================
# 批量导入
# =====================================================================

# 列名别名（中文表头 → 标准字段）
_COL_ALIASES = {
    "date": ["date", "交易日期", "日期", "date"],
    "direction": ["direction", "方向", "类型"],
    "amount": ["amount", "金额", "交易金额"],
    "account": ["account", "本方账户", "账户"],
    "dest_account": ["dest_account", "对方账户", "目标账户"],
    "category": ["category", "分类", "类别"],
    "tags": ["tags", "标签"],
    "description": ["description", "描述", "摘要"],
    "remark": ["remark", "备注"],
    "source_memo": ["source_memo", "原始摘要"],
    "counterparty": ["counterparty", "对方户名"],
    "location": ["location", "交易地点", "交易地点/附言", "附言"],
    "confidence": ["confidence", "置信度"],
}


def _build_col_index(headers):
    """headers: list[str]。返回 {标准字段: 列索引}。"""
    norm = {str(h).strip().lower(): i for i, h in enumerate(headers)}
    mapping = {}
    for std, aliases in _COL_ALIASES.items():
        for a in aliases:
            if a.lower() in norm:
                mapping[std] = norm[a.lower()]
                break
    return mapping


def _parse_amount(v):
    if isinstance(v, (int, float)):
        return float(v)
    if v is None:
        return 0.0
    s = str(v).strip().replace(",", "").replace("，", "")
    if s.startswith("(") and s.endswith(")"):
        s = "-" + s[1:-1]
    try:
        return float(s)
    except ValueError:
        return 0.0


def _read_file_rows(content: bytes, filename: str):
    """返回 (headers, rows)。rows 是 list[list[str]]。"""
    name = (filename or "").lower()
    if name.endswith(".csv"):
        # 尝试多种编码
        text = None
        for enc in ("utf-8-sig", "utf-8", "gbk", "gb18030"):
            try:
                text = content.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            text = content.decode("utf-8", errors="replace")
        reader = csv.reader(io.StringIO(text))
        rows_raw = list(reader)
    else:
        # xlsx/xls
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        rows_raw = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()

    if not rows_raw:
        return [], []
    return rows_raw[0], rows_raw[1:]


async def _resolve_user_maps(db: AsyncSession, user_id: str):
    """返回 (accounts_by_name, exp_cat_by_name, inc_cat_by_name, account_names, exp_cat_names, inc_cat_names)"""
    accs = (await db.execute(select(Account).where(Account.user_id == user_id))).scalars().all()
    cats = (await db.execute(select(Category).where(Category.user_id == user_id))).scalars().all()
    accounts_by_name = {a.name: a.id for a in accs}
    exp_cat = {c.name: c.id for c in cats if c.type == "expense"}
    inc_cat = {c.name: c.id for c in cats if c.type == "income"}
    return (accounts_by_name, exp_cat, inc_cat,
            [a.name for a in accs], list(exp_cat.keys()), list(inc_cat.keys()))


def _default_transfer_pair(accounts_by_name: dict, account: str) -> tuple[str, str]:
    """转账默认方向：工资账户 → 消费账户（BOOKKEEPING 规范）。

    返回 (转出账户, 转入账户)：
    - 本方是工资账户 → (工资, 消费)
    - 本方是消费账户且有工资账户 → (工资, 消费)，即把消费账户当转入方
    - 其他 → (本方, 第一个其他账户)
    """
    names = list(accounts_by_name.keys())
    others = [n for n in names if n != account]
    if not others:
        return (account, "")
    salary = next((n for n in names if "工资" in n), None)
    consumer = next((n for n in names if "消费" in n), None)
    if "工资" in account:
        return (account, consumer or others[0])
    if "消费" in account and salary:
        return (salary, account)
    return (account, others[0])


@router.post("/parse-import", response_model=ImportPreviewResponse)
async def parse_import(
    file: UploadFile = File(...),
    user_id: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_private_db),
):
    """解析导入文件，返回每行解析结果 + 重复检测，不写库。"""
    content = await file.read()
    headers, raw_rows = _read_file_rows(content, file.filename or "")
    if not headers:
        raise HTTPException(status_code=400, detail="文件为空或格式不支持")

    col = _build_col_index(headers)
    required = ["date", "direction", "amount", "account"]
    missing = [k for k in required if k not in col]
    if missing:
        raise HTTPException(status_code=400,
                            detail=f"缺少必需列: {missing}。请检查表头（需要：交易日期/方向/金额/本方账户）")

    accounts_by_name, exp_cat, inc_cat, acc_names, exp_names, inc_names = await _resolve_user_maps(db, user_id)

    # 查重：取该用户全部现有交易（含账户名，供前端展示重复候选）
    dup_result = await db.execute(
        select(Transaction, Account.name).join(Account, Account.id == Transaction.account_id).where(
            Transaction.user_id == user_id
        )
    )
    existing_by_key = {}
    for txn, acc_name in dup_result.all():
        key = (txn.account_id, txn.date, round(float(txn.amount), 2), txn.type)
        existing_by_key.setdefault(key, []).append({
            "date": txn.date,
            "amount": round(float(txn.amount), 2),
            "type": txn.type,
            "account": acc_name,
            "description": txn.description or "",
            "location": txn.location or "",
        })

    preview_rows = []
    err_count = 0
    dup_count = 0

    for idx, raw in enumerate(raw_rows, start=2):  # 行号从 2 开始（1=表头）
        def get(key):
            i = col.get(key)
            return (str(raw[i]).strip() if i is not None and i < len(raw) and raw[i] is not None else "")

        date = get("date")
        direction = get("direction").lower()
        amount = _parse_amount(get("amount"))  # 保留符号：负数=退款冲销（负支出）
        account = get("account")
        dest_account = get("dest_account")
        category = get("category")
        tags = get("tags")
        description = get("description")
        remark = get("remark")
        source_memo = get("source_memo")
        counterparty = get("counterparty")
        location = get("location")
        confidence = get("confidence")

        error = ""
        if direction not in ("income", "expense", "transfer"):
            error = f"方向无效: {direction!r}"
        elif amount == 0:
            error = "金额不能为 0"
        elif amount < 0 and direction not in ("expense", "income"):
            error = "负数金额仅支持支出（退款冲销）或收入（人工改判）"
        elif not re.match(r"^\d{4}-\d{2}-\d{2}", date):
            error = f"日期格式错误: {date!r}（需 YYYY-MM-DD）"
        elif account not in accounts_by_name:
            error = f"账户不存在: {account!r}"

        if not error and direction != "transfer" and category:
            cat_map = exp_cat if direction == "expense" else inc_cat
            if category not in cat_map:
                error = f"分类不存在: {category!r}（{direction} 分类）"

        if not error and direction == "transfer" and dest_account and dest_account not in accounts_by_name:
            error = f"对方账户不存在: {dest_account!r}"

        if not error and direction == "transfer" and not dest_account:
            src, dst = _default_transfer_pair(accounts_by_name, account)
            account, dest_account = src, dst

        # 查重
        is_dup = False
        dup_with = []
        if not error and account in accounts_by_name:
            key = (accounts_by_name[account], date[:10], round(amount, 2), direction)
            dup_with = existing_by_key.get(key, [])
            is_dup = bool(dup_with)

        if error:
            err_count += 1
        if is_dup:
            dup_count += 1

        preview_rows.append(ImportPreviewRow(
            row_index=idx, date=date, direction=direction, amount=amount,
            account=account, dest_account=dest_account, category=category,
            tags=tags, description=description, remark=remark,
            source_memo=source_memo, counterparty=counterparty,
            location=location, confidence=confidence,
            error=error, is_duplicate=is_dup,
            duplicate_with=dup_with[:5],
        ))

    return ImportPreviewResponse(
        rows=preview_rows, total=len(preview_rows),
        error_count=err_count, duplicate_count=dup_count,
        accounts=acc_names, expense_categories=exp_names, income_categories=inc_names,
    )


@router.post("/import", response_model=ImportResult)
async def commit_import(req: ImportCommitRequest, user_id: str = Depends(get_current_user_id), db: AsyncSession = Depends(get_private_db)):
    """按勾选的行批量入库。前端已做查重预览，这里再次校验账户/分类存在性。"""
    if not req.rows:
        return ImportResult(inserted=0, skipped=0, errors=[])

    accounts_by_name, exp_cat, inc_cat, _, _, _ = await _resolve_user_maps(db, user_id)

    inserted = 0
    skipped = 0
    errors = []

    for i, r in enumerate(req.rows, start=1):
        if r.direction not in ("income", "expense", "transfer"):
            errors.append(f"第{i}行: 方向无效 {r.direction!r}")
            skipped += 1
            continue
        if r.amount == 0:
            errors.append(f"第{i}行: 金额不能为 0")
            skipped += 1
            continue
        if r.amount < 0 and r.direction not in ("expense", "income"):
            errors.append(f"第{i}行: 负数金额仅支持支出（退款冲销）或收入（人工改判）")
            skipped += 1
            continue
        acc_id = accounts_by_name.get(r.account)
        if not acc_id:
            errors.append(f"第{i}行: 账户不存在 {r.account!r}")
            skipped += 1
            continue

        dest_id = None
        cat_id = None

        if r.direction == "transfer":
            if r.dest_account:
                dest_account_name = r.dest_account
            else:
                src_name, dest_account_name = _default_transfer_pair(accounts_by_name, r.account)
                if src_name != r.account:
                    acc_id = accounts_by_name.get(src_name)
                    if not acc_id:
                        errors.append(f"第{i}行: 账户不存在 {src_name!r}")
                        skipped += 1
                        continue
            if dest_account_name:
                dest_id = accounts_by_name.get(dest_account_name)
                if not dest_id:
                    errors.append(f"第{i}行: 对方账户不存在 {dest_account_name!r}")
                    skipped += 1
                    continue
        else:
            if r.category:
                cat_map = exp_cat if r.direction == "expense" else inc_cat
                cat_id = cat_map.get(r.category)
                if not cat_id:
                    errors.append(f"第{i}行: 分类不存在 {r.category!r}")
                    skipped += 1
                    continue

        tag_list = [t.strip() for t in r.tags.replace(";", ",").replace("；", ",").split(",") if t.strip()] if r.tags else []

        txn = Transaction(
            user_id=user_id, type=r.direction, date=r.date[:10],
            amount=(abs(r.amount) if r.direction == "income" else r.amount),
            account_id=acc_id, dest_account_id=dest_id, category_id=cat_id,
            tag_ids=tag_list, description=r.description or "", remark=r.remark or "",
            location=r.location or "",
        )
        db.add(txn)
        inserted += 1

    await db.commit()
    return ImportResult(inserted=inserted, skipped=skipped, errors=errors)